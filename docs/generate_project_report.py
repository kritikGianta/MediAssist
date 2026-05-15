from __future__ import annotations

from datetime import datetime
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "docs" / "MediAssist_AI_Project_Report.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_document_defaults(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    for section in document.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.85)
        section.right_margin = Inches(0.85)


def add_title_page(document: Document) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MediAssist AI\n")
    r.bold = True
    r.font.size = Pt(24)
    r.font.color.rgb = RGBColor(0x14, 0x26, 0x2F)

    r = p.add_run("Intelligent Healthcare Guidance Chatbot\n")
    r.bold = True
    r.font.size = Pt(18)

    p.add_run("\n")
    r = p.add_run("Comprehensive Project Documentation Report\n")
    r.font.size = Pt(14)

    p.add_run("\n")
    r = p.add_run(
        "Prepared for local-first development, testing, and deployment using only free/open-source tools.\n"
    )
    r.font.size = Pt(11)

    p.add_run("\n")
    r = p.add_run(f"Generated on: {datetime.now().strftime('%d %B %Y, %I:%M %p')}")
    r.italic = True

    p2 = document.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.add_run(
        "\nMedical Disclaimer: This chatbot is not a licensed medical professional. "
        "Consult a doctor for emergencies or accurate diagnosis."
    ).bold = True

    document.add_page_break()


def add_heading(document: Document, text: str, level: int = 1) -> None:
    document.add_heading(text, level=level)


def add_paragraph(document: Document, text: str) -> None:
    document.add_paragraph(text)


def add_bullets(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Bullet")


def add_numbered(document: Document, items: list[str]) -> None:
    for item in items:
        document.add_paragraph(item, style="List Number")


def add_table(document: Document, rows: list[list[str]], header_fill: str = "DDEFEA") -> None:
    table = document.add_table(rows=1, cols=len(rows[0]))
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for idx, value in enumerate(rows[0]):
        hdr[idx].text = value
        set_cell_shading(hdr[idx], header_fill)
    for row in rows[1:]:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    document.add_paragraph()


def build_report() -> None:
    document = Document()
    set_document_defaults(document)
    add_title_page(document)

    add_heading(document, "1. Executive Summary")
    add_paragraph(
        document,
        "MediAssist AI is a local-first healthcare guidance chatbot built with free and open-source tools. "
        "The system is designed to help users describe symptoms in natural language, receive general healthcare guidance, "
        "review uploaded health reports, explore non-prescription supportive medicine suggestions, and find nearby hospitals, "
        "doctors, and pharmacies. The project combines a modern React frontend with a Python FastAPI backend and a "
        "Retrieval-Augmented Generation (RAG) pipeline powered by FAISS, Sentence Transformers, and an Ollama-hosted local language model."
    )
    add_paragraph(
        document,
        "The main aim of the project is to create a healthcare support assistant that remains lightweight enough to run on a normal laptop, "
        "while still demonstrating modern AI architecture, safe design practices, real-world usability, and a beginner-friendly codebase."
    )

    add_heading(document, "2. Project Goals")
    add_bullets(
        document,
        [
            "Accept symptom descriptions in natural language.",
            "Retrieve relevant healthcare context using RAG before answering.",
            "Explain possible conditions in simple language.",
            "Provide safe general OTC/common-support suggestions only.",
            "Detect emergency warning patterns such as chest pain or breathing difficulty.",
            "Allow PDF upload for report summarization.",
            "Use OpenStreetMap-powered nearby healthcare search.",
            "Provide a premium-looking user interface suitable for a modern healthcare product demo.",
            "Use only free and open-source tooling.",
        ],
    )

    add_heading(document, "3. Problem Statement")
    add_paragraph(
        document,
        "Healthcare information is often fragmented across generic search engines, static informational websites, hospital directories, "
        "and clinical PDF reports. Many users need a single interface where they can ask symptom-based questions, review a health document, "
        "and decide whether they need nearby medical care. A plain conversational chatbot without retrieval may hallucinate or respond too generically. "
        "Therefore, this project uses Retrieval-Augmented Generation to ground answers in a healthcare knowledge base before generating a response."
    )

    add_heading(document, "4. What the System Does")
    add_bullets(
        document,
        [
            "Symptom checker: Users can type messages like 'I have fever, headache, and sore throat'.",
            "Possible condition hints: The assistant identifies likely conditions from retrieved context.",
            "Severity labeling: Responses are categorized as Mild, Moderate, or Serious.",
            "Precautions and self-care advice: The assistant suggests hydration, rest, monitoring, and similar general precautions.",
            "OTC support guidance: Only general non-dangerous, label-based medicine suggestions are allowed.",
            "Emergency escalation: The assistant explicitly warns users to seek immediate medical attention when severe symptoms are detected.",
            "PDF summarization: Users can upload a health-related PDF and get a text summary.",
            "Nearby doctor/hospital finder: Uses geolocation and OpenStreetMap/Overpass API to list nearby facilities.",
            "Persistent settings and history: The system stores chat history and model settings locally.",
        ],
    )

    add_heading(document, "5. Why Use LLMs in This Project")
    add_paragraph(
        document,
        "Large Language Models (LLMs) are useful in this project because healthcare users typically describe their symptoms in natural, conversational language rather than structured forms. "
        "A traditional rule-based system can recognize a fixed symptom list, but it struggles with free-form input, paraphrased descriptions, incomplete context, or follow-up questions. "
        "An LLM can interpret natural language, maintain conversational continuity, synthesize retrieved medical context, and return a readable explanation."
    )
    add_paragraph(
        document,
        "However, LLMs alone are not enough for trustworthy healthcare guidance. A pure model-only chatbot may produce unsupported or hallucinated answers. "
        "That is why this system does not rely on a general LLM alone; it combines a local LLM with a RAG retrieval layer and rule-based medical guardrails."
    )

    add_heading(document, "6. Why Use a Local Open-Source LLM")
    add_bullets(
        document,
        [
            "Cost control: Running a local open-source model avoids paid API usage.",
            "Offline/local-first capability: The assistant can function locally without depending on proprietary cloud inference.",
            "Privacy improvement: Sensitive symptom descriptions can remain on the user’s machine or local network environment.",
            "Educational value: The project demonstrates how modern LLM systems can be built end-to-end with open tools.",
            "Customization: Local models can be swapped or tuned later.",
            "Compliance with the project requirement: The user explicitly requested only free/open-source tools.",
        ],
    )
    add_paragraph(
        document,
        "The default intended model path is Ollama with a lightweight model such as TinyLlama. The report also notes that in the current local setup, "
        "if Ollama is not installed or not running, the backend falls back to a deterministic safe response path so the rest of the application remains testable."
    )

    add_heading(document, "7. Why Retrieval-Augmented Generation (RAG) Was Chosen")
    add_paragraph(
        document,
        "RAG is used to improve answer grounding. Instead of asking the model to answer purely from memorized parameters, the system first retrieves the most relevant healthcare text chunks from a local knowledge base. "
        "Those chunks are then passed to the LLM along with the user query. This produces more context-aware and evidence-backed answers."
    )
    add_bullets(
        document,
        [
            "Improves reliability compared to model-only answering.",
            "Allows domain customization using healthcare-specific text sources.",
            "Makes it easier to update the knowledge base without retraining the model.",
            "Supports source citation in the UI.",
            "Works well on a laptop because embeddings + FAISS are relatively lightweight for a small knowledge base.",
        ],
    )

    add_heading(document, "8. Core Architecture")
    add_table(
        document,
        [
            ["Layer", "Technology", "Purpose"],
            ["Frontend", "React + Vite + Tailwind CSS + Framer Motion", "User interface, animations, chat, upload, nearby care, settings"],
            ["Backend", "FastAPI", "REST API endpoints and application orchestration"],
            ["RAG Retrieval", "LangChain + FAISS", "Document chunking, embeddings, similarity search"],
            ["Embeddings", "Sentence Transformers / all-MiniLM-L6-v2", "Vector representation of text chunks"],
            ["Local LLM", "Ollama + TinyLlama/Phi/Gemma/FLAN-T5 path", "Conversational answer generation"],
            ["PDF Handling", "PyPDF", "Extract text from uploaded health reports"],
            ["Nearby Search", "OpenStreetMap Overpass API", "Locate nearby hospitals, clinics, doctors, pharmacies"],
            ["Persistence", "JSON files", "Chat history and settings storage"],
        ],
    )

    add_heading(document, "9. Frontend Design and User Experience")
    add_paragraph(
        document,
        "The frontend was designed to move away from a generic 'AI demo' appearance and towards a softer healthcare-product style. "
        "The application now includes a dedicated landing page with clear product messaging, a call-to-action button to open the chatbot workspace, "
        "and a workspace page that focuses on interaction instead of marketing content."
    )
    add_bullets(
        document,
        [
            "Landing page: Hero section with summary of system capabilities and a bot illustration.",
            "Workspace page: Dedicated area for chat, report upload, experimentation panel, and nearby care search.",
            "Responsive design: Uses Tailwind utility classes and adaptive grid layouts.",
            "Motion: Framer Motion is used for subtle transitions rather than excessive visual noise.",
            "Voice UX: Browser speech recognition and speech synthesis are used when available.",
        ],
    )

    add_heading(document, "10. Backend Design")
    add_paragraph(
        document,
        "The backend is organized into small modules for configuration, models, utilities, and services. "
        "This keeps the project beginner-friendly while still demonstrating clean separation of responsibilities."
    )
    add_table(
        document,
        [
            ["Module/File", "Responsibility"],
            ["app/main.py", "FastAPI app setup and route registration"],
            ["app/models.py", "Pydantic request/response schemas"],
            ["app/config.py", "Default settings, file paths, runtime initialization"],
            ["app/services/chat_service.py", "Chat orchestration, prompt building, history, safe fallback"],
            ["app/services/rag_service.py", "Load knowledge data, split chunks, embed text, FAISS retrieval"],
            ["app/services/location_service.py", "Nearby healthcare search using Overpass API"],
            ["app/services/report_service.py", "PDF text extraction and report summarization"],
            ["app/utils/medical_guardrails.py", "Emergency keyword detection, severity inference, safe OTC suggestions"],
            ["app/utils/storage.py", "Local JSON read/write helpers"],
        ],
    )

    add_heading(document, "11. RAG Pipeline: How It Works")
    add_numbered(
        document,
        [
            "Healthcare documents are loaded from CSV and optional text files.",
            "The content is split into chunks using RecursiveCharacterTextSplitter.",
            "Each chunk is converted into embeddings using all-MiniLM-L6-v2.",
            "Embeddings are stored in a FAISS vector index.",
            "When the user asks a question, the system performs similarity search over the indexed chunks.",
            "Top-k relevant chunks are selected as context.",
            "The prompt includes the user message, recent history, severity hints, emergency flags, and retrieved context.",
            "The local Ollama model generates a grounded response.",
            "If the model is unavailable, a deterministic fallback answer is returned using rule-based safety logic.",
        ],
    )

    add_heading(document, "12. Medical Safety Design")
    add_paragraph(
        document,
        "Healthcare applications must be especially careful about safety, scope, and overconfidence. "
        "This system explicitly avoids positioning itself as a diagnostic authority and instead acts as a guidance assistant."
    )
    add_bullets(
        document,
        [
            "Mandatory disclaimer is included in responses and health endpoint text.",
            "Emergency keywords such as chest pain, breathing difficulty, stroke signs, seizures, or fainting trigger urgent warnings.",
            "Only general OTC/common-support suggestions are allowed.",
            "No dangerous prescriptions, opioids, antibiotics, or steroid recommendations are provided.",
            "Users are encouraged to consult a doctor for severe, worsening, or unclear symptoms.",
            "The report summarization feature avoids pretending to interpret clinical values as a doctor.",
        ],
    )

    add_heading(document, "13. Datasets and Knowledge Sources")
    add_paragraph(
        document,
        "The system is structured to support real public healthcare sources. Because the user requested a storage-efficient local build, the current repository includes a compact sample healthcare dataset. "
        "The architecture is intentionally prepared so larger real datasets can be added later when storage allows."
    )
    add_table(
        document,
        [
            ["Source", "Role in Project", "Status in Current Build"],
            ["MedQuAD", "Medical Q&A retrieval source", "Referenced for future ingestion"],
            ["Disease Symptom Description Dataset", "Symptoms, conditions, precautions", "Referenced for future ingestion"],
            ["openFDA", "OTC/drug labeling concepts", "Referenced conceptually and for extension"],
            ["OpenStreetMap / Overpass", "Nearby care discovery", "Actively used"],
            ["Local healthcare_knowledge.csv", "Small sample retrieval knowledge base", "Actively used"],
        ],
    )

    add_heading(document, "14. API Endpoints")
    add_table(
        document,
        [
            ["Endpoint", "Method", "Purpose"],
            ["/health", "GET", "Backend health status and disclaimer"],
            ["/", "GET", "Basic backend status page"],
            ["/chat", "POST", "Main symptom-guidance/chat endpoint"],
            ["/upload-report", "POST", "Upload PDF and return summary"],
            ["/nearby-doctors", "POST", "Nearby hospitals/clinics/doctors/pharmacies search"],
            ["/settings", "GET/POST", "Read and save runtime generation settings"],
            ["/history", "GET", "Return locally stored chat history"],
        ],
    )

    add_heading(document, "15. Why the System Uses FastAPI")
    add_bullets(
        document,
        [
            "FastAPI is free and open-source.",
            "It provides automatic request validation through Pydantic models.",
            "It is beginner-friendly yet production-capable.",
            "It offers fast development speed for AI APIs.",
            "Interactive API docs are available through Swagger/OpenAPI support.",
        ],
    )

    add_heading(document, "16. Why FAISS and Sentence Transformers Were Chosen")
    add_paragraph(
        document,
        "FAISS is a highly popular open-source similarity search library that is efficient for vector retrieval tasks. "
        "Sentence Transformers provides high-quality text embeddings with small and practical models such as all-MiniLM-L6-v2, which is well suited to low-storage and local-laptop constraints."
    )
    add_bullets(
        document,
        [
            "Small embedding model footprint compared to many alternatives.",
            "Good semantic retrieval quality for short healthcare text chunks.",
            "Fast enough for local development and experimentation.",
            "Straightforward integration with LangChain.",
        ],
    )

    add_heading(document, "17. Why Ollama Was Chosen")
    add_paragraph(
        document,
        "Ollama simplifies local LLM execution by handling model download, local serving, and inference management behind a very simple interface. "
        "It fits the educational and local-first goals of the project. Instead of requiring the user to manage a complicated model runtime manually, Ollama offers a clean developer workflow."
    )
    add_bullets(
        document,
        [
            "Simple local model serving command interface.",
            "Works well with open-source models.",
            "Integrates easily with LangChain community wrappers.",
            "Supports swapping small models based on hardware capacity.",
        ],
    )

    add_heading(document, "18. Current Limitations")
    add_bullets(
        document,
        [
            "The current repository uses a compact sample healthcare knowledge base rather than full MedQuAD ingestion.",
            "Clinical report analysis is text-summary based and does not perform specialist-grade lab interpretation.",
            "Nearby search quality depends on OpenStreetMap data coverage.",
            "If Ollama is not installed or running, the system uses a fallback safe-answer mode instead of true local LLM inference.",
            "The current setup does not include user authentication or encrypted long-term health record storage.",
            "This assistant is not a substitute for professional diagnosis.",
        ],
    )

    add_heading(document, "19. Testing and Verification Performed")
    add_bullets(
        document,
        [
            "Frontend dependency installation and production build validation with Vite.",
            "Backend dependency installation and Python compile validation.",
            "Health endpoint verification on the FastAPI backend.",
            "Settings and history endpoint verification.",
            "Nearby healthcare search testing path enabled through the frontend and API.",
            "UI iteration and multiple layout corrections on the landing page and chat workspace.",
        ],
    )

    add_heading(document, "20. Deployment Strategy")
    add_paragraph(
        document,
        "The project is designed for local-first execution and later staged deployment. "
        "The frontend can be deployed on Vercel free tier and the backend shell can be deployed on Render free tier. "
        "However, true local-model inference through Ollama is generally better suited to local execution or a compatible private server rather than a limited free cloud instance."
    )
    add_bullets(
        document,
        [
            "Frontend target: Vercel free tier.",
            "Backend target: Render free tier for API shell.",
            "Recommended inference mode for real use: Local Ollama or a self-managed server with local model support.",
        ],
    )

    add_heading(document, "21. Future Enhancements")
    add_bullets(
        document,
        [
            "Add a real ingestion script for MedQuAD and disease-symptom datasets.",
            "Support structured CSV/PDF ingestion pipelines from more healthcare documents.",
            "Migrate to newer LangChain provider packages such as langchain-ollama and langchain-huggingface.",
            "Add authentication and role-based access for clinician or patient modes.",
            "Improve PDF lab result extraction and abnormal-value highlighting.",
            "Add maps and specialist filtering in nearby search.",
            "Add multilingual support for wider accessibility.",
            "Add unit tests and integration tests for chat, upload, and retrieval flows.",
        ],
    )

    add_heading(document, "22. Conclusion")
    add_paragraph(
        document,
        "MediAssist AI demonstrates how a meaningful healthcare guidance assistant can be built using only free/open-source tools. "
        "Its value comes from combining a local open-source LLM with retrieval, structured safeguards, practical document handling, and real-world nearby-care lookup. "
        "The project is intentionally designed to be understandable, extendable, and feasible for student or portfolio use while still reflecting modern AI application architecture."
    )
    add_paragraph(
        document,
        "In summary, this project is not just a chatbot. It is a healthcare support workflow that joins conversational AI, local inference, RAG retrieval, UI design, medical disclaimers, and geolocation-based care discovery into one cohesive system."
    )

    document.add_section(WD_SECTION.NEW_PAGE)
    add_heading(document, "Appendix A: Key Local File References")
    add_bullets(
        document,
        [
            str(ROOT / "backend" / "app" / "main.py"),
            str(ROOT / "backend" / "app" / "services" / "chat_service.py"),
            str(ROOT / "backend" / "app" / "services" / "rag_service.py"),
            str(ROOT / "backend" / "app" / "services" / "location_service.py"),
            str(ROOT / "frontend" / "src" / "App.jsx"),
            str(ROOT / "frontend" / "src" / "components" / "LandingPage.jsx"),
            str(ROOT / "frontend" / "src" / "components" / "Sidebar.jsx"),
            str(ROOT / "README.md"),
        ],
    )

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(str(OUTPUT))
    print(f"Report written to: {OUTPUT}")


if __name__ == "__main__":
    build_report()
