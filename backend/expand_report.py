import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

def expand_report(file_path):
    doc = Document(file_path)
    new_doc = Document()
    
    for para in doc.paragraphs:
        new_para = new_doc.add_paragraph()
        new_para.alignment = para.alignment
        new_para.paragraph_format.space_after = para.paragraph_format.space_after
        new_para.paragraph_format.line_spacing = para.paragraph_format.line_spacing
        new_para.paragraph_format.left_indent = para.paragraph_format.left_indent
        new_para.style = para.style
        
        for run in para.runs:
            new_run = new_para.add_run(run.text)
            new_run.bold = run.bold
            new_run.italic = run.italic
            new_run.font.name = run.font.name
            new_run.font.size = run.font.size
            if run.font.color and run.font.color.rgb:
                new_run.font.color.rgb = run.font.color.rgb
            
        if "11. Generative AI & RAG Pipeline Architecture" in para.text:
            p_flow = new_doc.add_paragraph()
            p_flow.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_flow = p_flow.add_run("\nSystem Architecture Workflow Diagram\n")
            r_flow.bold = True
            r_flow.font.name = 'Times New Roman'
            r_flow.font.size = Pt(14)
            
            diagram = (
                "[ User Input (React Frontend) ]\n"
                "        │\n"
                "        ▼\n"
                "[ FastAPI Backend: Chat Service ]\n"
                "        │\n"
                "        ├─► (1) Intent & Safety Check (Emergency/Severity detection)\n"
                "        │\n"
                "        ▼\n"
                "[ RAG Retrieval Service ]\n"
                "        ├─► Reads from FAISS Vector Store (SentenceTransformers Embeddings)\n"
                "        ├─► Matches with MedlinePlus (NIH) & Local CSV Knowledge\n"
                "        │\n"
                "        ▼\n"
                "[ Advanced Prompt Engineering ]\n"
                "        ├─► Injects: User Query + Top-k Medical Context + Chat History + Guardrails\n"
                "        │\n"
                "        ▼\n"
                "[ Generative LLM (Groq API) ]\n"
                "        ├─► Llama 3.1 8B/70B evaluates context and synthesizes response\n"
                "        │\n"
                "        ▼\n"
                "[ Final Response Generation ]\n"
                "        ├─► Attaches Citations, Severity Warnings, and Medical Disclaimer\n"
                "        │\n"
                "        ▼\n"
                "[ User Output (React Frontend) ]\n"
            )
            p_diag = new_doc.add_paragraph()
            p_diag.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r_diag = p_diag.add_run(diagram)
            r_diag.font.name = 'Courier New'
            r_diag.font.size = Pt(10)
            
            p_exp = new_doc.add_paragraph()
            p_exp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            r_exp = p_exp.add_run("Workflow Explained Fast:\n")
            r_exp.bold = True
            r_exp.font.name = 'Times New Roman'
            r_exp.font.size = Pt(12)
            
            p_steps = new_doc.add_paragraph()
            p_steps.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p_steps.paragraph_format.line_spacing = 1.5
            r_steps = p_steps.add_run(
                "1. Data Ingestion: Medical literature from MedlinePlus is chunked and stored as embeddings in a FAISS vector database.\n"
                "2. User Query: The user asks a health-related question via the React frontend.\n"
                "3. Safety Check: The backend immediately scans for emergency keywords to prioritize urgent warnings.\n"
                "4. Semantic Search: The RAG engine searches the FAISS database for the most relevant medical facts.\n"
                "5. LLM Synthesis: The Groq API (using Llama 3.1) receives the prompt containing the context and generates a safe, human-like response.\n"
                "6. Delivery: The frontend displays the formatted answer along with proper medical citations and disclaimers."
            )
            r_steps.font.name = 'Times New Roman'
            r_steps.font.size = Pt(12)
            
    new_doc.save(file_path)
    print("Report expanded successfully.")

if __name__ == "__main__":
    expand_report(r"c:\Users\kriti\Downloads\projects_final\medical chatbot\docs\MediAssist_AI_Project_Report.docx")
